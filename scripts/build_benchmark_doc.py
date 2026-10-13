# ruff: noqa: E501
"Build editable upgrade notes and append them to a copy of the original introduction."

from __future__ import annotations

import json
from pathlib import Path

from build_submission_doc import configure, editable_table, paragraph
from docx import Document
from docx.shared import Cm

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/contest/submission"


def notes(doc):
    doc.add_heading("题库与可复现评测升级", 0)
    paragraph(doc, "麟阁 MemHall · 2026 年 10 月 14 日 · 参赛作品补充说明", color="607286")
    paragraph(
        doc,
        "本轮补强数据设计、自动评分和跨智能体公平比较。原 49 道核心题与历史成绩保留原口径，新增挑战题的真实智能体结果待单独实验。",
    )
    doc.add_heading("1. 从题量增加到场景覆盖", 1)
    paragraph(
        doc,
        (
            "长期记忆的难点包括旧值替换、局部更新、相近实体、不同作用域、撤销与不可信引用，以及把记住的格式用于后续真实任务。本轮将这"
            "些情形拆成可追踪的场景模板，而不是仅对原有路径题换词。"
        ),
    )
    editable_table(
        doc,
        ["题包", "实例", "显式模板", "用途"],
        [
            ["full + chains", "49", "原题未声明", "历史核心口径"],
            ["challenge-v2", "72", "24", "公开测试：每模板 3 变体"],
            ["development-v2", "39", "13", "开发集：与测试模板分离，含 canary 读侧探针试点"],
            ["生成测试池示例", "1008", "24", "每模板 42 变体，非 1008 独立题族"],
        ],
    )
    paragraph(
        doc,
        (
            "新增 37 个模板覆盖六维能力。开发和测试按模板分开；两个集合均公开，不能声称是私有 heldout。变体记录模板版本、"
            "内容哈希、seed、序号、干扰轮数及取值参数，支持重新生成和内容对账。"
        ),
    )
    paragraph(
        doc,
        (
            "难度 1／2／3 为设计标签，默认配 1／4／12 轮干扰，尚待真实通过率校准。新增题通过切换会话测跨会话保留，不冒充真"
            "实等待数天、系统重启或真实工作历史。模板数量同样不自动证明统计独立性。"
        ),
    )

    doc.add_page_break()
    doc.add_heading("2. 让不同智能体做同一份试卷", 1)
    paragraph(
        doc,
        (
            "历史轮换读取过去的评测频次，相同 seed 在不同历史下也会换题。本轮新增分层抽样：只根据题包与明确参数选题，封存完整题"
            "目后再交给各智能体。"
        ),
    )
    editable_table(
        doc,
        ["封存项", "作用"],
        [
            ["题包 ID、版本、内容指纹", "防止同名题库在更新后悄悄改变"],
            ["完整用例、题单 ID、实际清单", "换机仍可运行，内容改动会被拒绝"],
            ["seed、能力／内容／难度分布", "说明抽样过程与覆盖"],
            ["评分协议与判卷器配置", "避免不同判卷条件混在同一榜"],
        ],
    )
    for command in [
        (
            "python -m memhall sample -c cases/challenge-v2 --count 24 --"
            "seed 20261013 --max-per-family 1 -o build/paper-24.json"
        ),
        "python -m memhall run -a mock --paper build/paper-24.json",
        (
            "python -m memhall batch --paper build/paper-24.json --adapte"
            "rs hermes kylinbot openclaw --repeats 2 --judge dual"
        ),
    ]:
        paragraph(doc, command, size=9, color="142C46")
    paragraph(
        doc,
        (
            "批量命令由操作者在配置好真实车道后启动，记录每轮目录和退出状态。试卷文件拒绝覆盖；封存后不能再做历史轮换。UI 支持生成"
            "、导入和导出固定试卷，并选择它开始评测。"
        ),
    )
    paragraph(
        doc,
        (
            "compare／aggregate 在输出前核对实际题目、题包、试卷、判卷协议和产物封印。相同题号配不同参数或不同评委配"
            "置会被拒绝。旧归档缺少封存信息时标明限制，新旧不可核对的协议不直接混用。"
        ),
    )

    doc.add_page_break()
    doc.add_heading("3. 数据、适配器与断言可独立扩展", 1)
    paragraph(
        doc,
        (
            "题包使用 YAML 与 suite.yaml，声明版本、许可、默认 split、评分协议和相对 glob。CLI 可直接"
            "加载外部目录；读取题包不会执行其中的 Python，目录外引用和重复题号／探测点会被拒绝。"
        ),
    )
    paragraph(
        doc,
        (
            "适配器继承 AgentAdapter，实现清空隔离状态、发消息、切换会话、导出记忆和操作记录。外部 Python 包通过"
            "入口点声明自己，不需要改核心注册表："
        ),
    )
    for line in [
        '[project.entry-points."memhall.adapters"]',
        'example = "my_memhall_plugin:ExampleAdapter"',
        '[project.entry-points."memhall.assertions"]',
        '"example.valid_record" = "my_memhall_plugin:valid_record"',
    ]:
        paragraph(doc, line, size=9)
    paragraph(
        doc,
        (
            "适配器和断言在显式选择时加载，禁止覆盖内置同名实现。断言接收参数及统一 EvidenceStore，返回布尔值；评分归档"
            "记录使用到的外部断言实现指纹。进程内也可显式注册。UI 的部署探测仍以既有注册表为准，外部插件可通过 CLI 名称使用。"
        ),
    )
    paragraph(
        doc,
        (
            "扩展题包可以只覆盖某一任务；缺失维度显示未测，不补零分。外部 QA 数据导入后必须保留原任务语义，不能宣称自动覆盖了 o"
            "penKylin 文件执行或敏感持久化。已安装插件需要审查其 reset 范围，并使用隔离工作区。"
        ),
    )
    paragraph(
        doc,
        (
            "完整入口契约、题包样板、CLI 命令和限制见仓库 docs/benchmark-v2.md。公开模板生成器只支持 dev"
            "elopment／test／stress，不把公开规则包装成私有测试集。"
        ),
    )

    doc.add_page_break()
    doc.add_heading("4. 任务复用由实际产物验收", 1)
    paragraph(
        doc,
        (
            "新复用题先教会格式并生成初始产物，切换会话、经过干扰后，再给项目和新的状态／目标标签，让智能体按记住的约定输出新文件。测"
            "试模板包括 CSV 报告、JSON 配置、发布清单和回滚计划。"
        ),
    )
    editable_table(
        doc,
        ["证据", "本轮验收规则"],
        [
            ["回复与输出路径", "诊断信息，不凭一句已完成给任务成功分"],
            ["文件内容 SHA-256", "采自被测环境，逐字节核对格式与字段"],
            ["创建／修改及阶段窗", "确认采样窗内出现实际产物变化"],
            ["内容证据缺失", "标记运行无效，不在判卷机本地猜测"],
        ],
    )
    paragraph(
        doc,
        (
            "fs.artifact_matches 同时检查路径、内容哈希与阶段创建／修改。SSH 车道在远端只读采样，直连车道仅读"
            "取自己的隔离工作区；内容上限为每文件 2 MiB，链接、不可读文件与超限文件不提供内容指纹。"
        ),
    )
    paragraph(
        doc,
        (
            "目前实现内容采样的车道包括 hermes、kylinbot、openclaw、hermes-local、claude-l"
            "ocal 及 mock。其他插件可以补充可选接口。mock 的空文件工作区会让只说不做的复用行为失败，而不是凭语言回答过"
            "关。"
        ),
    )
    paragraph(
        doc,
        (
            "阶段窗为 base→inject 结束和 inject 结束→最终采样，第二个窗包含 confound。哈希证明的是采样"
            "时点的文件符合指定格式，不能证明任意业务任务的所有中间步骤、外部副作用或业务效果。"
        ),
    )

    doc.add_page_break()
    doc.add_heading("5. 统计与判卷质量有独立出口", 1)
    paragraph(
        doc,
        (
            "总分保留 v2 的有效 score 探测点正确率。diagnostic、运行无效、判卷未决和降级判定不进正式分母。新增题"
            "族等权辅助分，不替代历史总分。"
        ),
    )
    paragraph(
        doc,
        (
            "置信区间把同一题族的探测点一起重采样，并保持整轮波动；配对方向检验先合并同题，再合并同族变体。旧题缺少 family_i"
            "d 时按 case_id 聚类，不代表它们已被证明独立。单题族区间和少量轮次不支撑强排名；单对检验也不自动校正多重比较。"
        ),
    )
    paragraph(
        doc,
        (
            "新增人工盲评工具按能力抽样，工作表隐藏智能体身份和机器结论，答案键另存。人工只填 human_verdict 和备注；问"
            "题内容、来源和答案键被修改会被拒绝。"
        ),
    )
    for command in [
        (
            "python -m memhall calibrate runs/某轮A runs/某轮B --per-capabili"
            "ty 12 --seed 20261013 -o build/human-audit"
        ),
        "python -m memhall calibrate --analyze build/human-audit",
    ]:
        paragraph(doc, command, size=9)
    paragraph(
        doc,
        (
            "完成真实人工标签后，可输出六态混淆矩阵、对错一致率、误判为正确的数量、分能力统计和 Cohen κ。空白保持待人工标签，"
            "部分填写保持 partial。离线脚本样本可验证工具，但不是实际 LLM 校准，也不是独立双评委实验。"
        ),
    )

    doc.add_page_break()
    doc.add_heading("6. 验证证据与仍待完成的实验", 1)
    validation_path = ROOT / "docs/audit/review-20261013/benchmark-upgrade-validation.json"
    validation = json.loads(validation_path.read_text(encoding="utf-8")) if validation_path.exists() else {}
    paragraph(
        doc,
        (
            "离线实际生成 1008 道实例，仍为 24 个模板。固定 24 题 mock 轮产出 168 条证据，通过完整性校验；每"
            "题只有一次 inject 记忆采样的提示保留。共 32 探测（24 score、8 diagnostic），22 个有效"
            " score、2 个待复核，有效总体为 54.55%。这是工具链演示和缺陷注入基线。"
        ),
    )
    if validation.get("pytest"):
        paragraph(
            doc, f"工程验证：{validation['pytest']}。Ruff、mypy、题库门禁与离线判卷自检结果随验证清单保存。"
        )
    paragraph(
        doc,
        (
            "回归覆盖抽样确定性、跨车道同卷、试卷搬移、内容与协议混用拒绝、产物篡改拒绝、题族聚类、真实临时文件改写、插件加载、人工盲"
            "评完整性及 UI／CLI。"
        ),
    )
    editable_table(
        doc,
        ["仍需补充", "原因"],
        [
            ["至少两款真实智能体的新题重复实验", "验证区分度、有效率与模板难度"],
            ["人工标签与独立评委 B 结果", "工程自检不能证明语义判卷准确"],
            ["干净 openKylin 上重建发行包与录屏", "源码门禁不能替代目标环境交付验收"],
            ["历史缺失原件或重跑归档", "原始哈希和历史成绩不改写"],
        ],
    )
    paragraph(
        doc,
        (
            "命题来源为 2026-10-13 核对的 openKylin 官方赛题，原文与 PDF 留存在 docs/audit/r"
            "eview-20261013/benchmark-sources；数据设计的前期比较与出处见 MemHall-题库与扩展"
            "设计评估.md。"
        ),
    )
    screenshot = OUT / "assets/ui-frozen-paper.png"
    if screenshot.exists():
        doc.add_picture(str(screenshot), width=Cm(15.5))
        paragraph(doc, "固定试卷 UI：本轮实际页面测试截图，mock 模式。", size=8, color="607286")


    doc.add_page_break()
    doc.add_heading("7. LLM 辅助判卷如何符合赛题", 1)
    paragraph(doc, "2026 年 10 月 14 日重新下载官方命题 PDF，内容哈希与审计留存版一致。公开文本要求自动化、低人工介入、可复现、可扩展，并依据运行证据给出可解释评分；没有限定判卷算法，也没有禁止 LLM。因此，我们判断 LLM 可以参与语义判卷。这是技术解读，不代表主办方的另行许可或背书。")
    paragraph(doc, "命题同时把自动评分的语义偏移列为痛点，自动评分能力占 25%，稳定性与可复现性占 15%。参赛方案需要证明判卷可靠性；模型自报置信度、理由看起来合理或两位评委意见一致，都不能独立证明结论正确。")
    editable_table(doc, ["判定对象", "采用的方式与验证出口"], [
        ["路径、文件内容和阶段状态", "确定性断言与真实环境采样，任务复用检查实际产物"],
        ["改写、遗漏、混淆等语义", "固定标准下的 LLM 结构化判定，附理由和证据引用"],
        ["语义判卷准确度与偏差", "人工分层盲评，报告一致率、误判正确率和未决率"],
        ["重复性与跨智能体公平性", "同卷、同协议与评委配置封存，重复判卷与交叉核对"],
    ])
    paragraph(doc, "当前历史 full/chains 实测只有评委 A，且 A 与被测智能体同一模型。--judge dual 是配置模式名称，B 必须实际配置并成功调用才构成双评委。人工盲评工具已实现，但真实人工标签和独立 B 实验尚未完成，材料如实披露此限制。")
    paragraph(doc, "证据不足、无效输出和未解决争议保留未决状态，输出有效覆盖率与保守下界。人工介入主要用于校准和异常复核；无需把全部题目交由人工判卷。官方来源、判卷配置及验证记录见 provenance.json、docs/benchmark-v2.md 与审计目录。")


def build():
    doc = Document()
    configure(doc)
    notes(doc)
    doc.core_properties.title = "MemHall 题库与可复现评测升级说明"
    doc.save(OUT / "MemHall-题库与评测升级说明.docx")
    introduction = Document(OUT / "MemHall-作品介绍-可编辑版.docx")
    introduction.add_page_break()
    notes(introduction)
    introduction.core_properties.title = "MemHall 作品介绍 · 题库升级版"
    introduction.save(OUT / "MemHall-作品介绍-题库升级版.docx")
    print("Saved editable upgrade notes and a copy of the introduction with the upgrade appendix.")


if __name__ == "__main__":
    build()
