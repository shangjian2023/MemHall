"""Build the editable contest document from public content and frozen experiment data.

Run: uv run --group docs python scripts/build_submission_doc.py
The original experimental runs are never modified by this builder.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from matplotlib.patches import Circle, Wedge

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/contest/submission"
ASSETS = OUT / "assets"
INK = "142C46"
TEAL = "087F86"
MUTED = "607286"
LIGHT = "F0F5F8"
ORDER = ["hermes", "openclaw", "kylinbot"]
COLORS = ["#138C8B", "#3869AA", "#B88735"]
CAPS = ["persist", "recall", "dynamic_update", "discriminate", "boundary", "reuse"]


def xml(tag, **attrs):
    element = OxmlElement(tag)
    for name, value in attrs.items():
        element.set(qn(name), str(value))
    return element


def font(run, size=None, bold=None, color=None):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().append(xml("w:rFonts", **{"w:eastAsia": "微软雅黑"}))
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def shade(cell, fill):
    cell._tc.get_or_add_tcPr().append(xml("w:shd", **{"w:fill": fill}))


def paragraph(container, text, style=None, size=None, color=None, bold=None):
    para = container.add_paragraph(style=style)
    font(para.add_run(text), size=size, color=color, bold=bold)
    return para


def field(para, instruction):
    run = para.add_run()
    run._r.append(xml("w:fldChar", **{"w:fldCharType": "begin"}))
    text = xml("w:instrText", **{"xml:space": "preserve"})
    text.text = instruction
    run._r.append(text)
    run._r.append(xml("w:fldChar", **{"w:fldCharType": "separate"}))
    font(para.add_run("目录将在 Word 中更新"), color=MUTED, size=10)
    para.add_run()._r.append(xml("w:fldChar", **{"w:fldCharType": "end"}))


def hyperlink(para, label, url):
    rel = para.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    link = xml("w:hyperlink", **{"r:id": rel})
    run = xml("w:r")
    props = xml("w:rPr")
    props.append(xml("w:color", **{"w:val": TEAL}))
    props.append(xml("w:u", **{"w:val": "single"}))
    run.append(props)
    text = xml("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    para._p.append(link)


def editable_table(doc, headers, rows, sizes=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    if sizes is None:
        ratios = {2: [0.28, 0.72], 3: [0.22, 0.39, 0.39]}.get(len(headers), [1 / len(headers)] * len(headers))
        sizes = [16.6 * x for x in ratios]
    for i, width in enumerate(sizes):
        table.columns[i].width = Cm(width)
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Cm(sizes[i])
        shade(cell, INK)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        para = cell.paragraphs[0]
        para.paragraph_format.space_after = Pt(5)
        para.paragraph_format.space_before = Pt(5)
        font(para.add_run(str(text)), size=9.2, bold=True, color="FFFFFF")
    table.rows[0]._tr.get_or_add_trPr().append(xml("w:tblHeader"))
    for n, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].width = Cm(sizes[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade(cells[i], LIGHT if n % 2 == 0 else "FFFFFF")
            para = cells[i].paragraphs[0]
            para.paragraph_format.space_before = Pt(5)
            para.paragraph_format.space_after = Pt(5)
            para.paragraph_format.line_spacing = 1.17
            font(para.add_run(str(text)), size=9.0, color=INK, bold=i == 0)
        table.rows[-1]._tr.get_or_add_trPr().append(xml("w:cantSplit"))
    table._tbl.tblPr.append(xml("w:tblBorders"))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def configure(doc):
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.top_margin = Cm(1.85)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.75)
    section.footer_distance = Cm(0.85)
    section.different_first_page_header_footer = True
    specs = {
        "Normal": (10.3, INK),
        "Heading 1": (19, INK),
        "Heading 2": (11.7, TEAL),
        "Caption": (8.7, MUTED),
    }
    for name, (size, color) in specs.items():
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.get_or_add_rPr().append(xml("w:rFonts", **{"w:eastAsia": "微软雅黑"}))
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.32
    for name in ["Heading 1", "Heading 2"]:
        doc.styles[name].font.bold = True
        doc.styles[name].paragraph_format.keep_with_next = True
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(12)
    doc.styles["Caption"].paragraph_format.space_after = Pt(9)
    header = section.header.paragraphs[0]
    font(header.add_run("麟阁 MEMHALL"), size=9.5, color=TEAL, bold=True)
    font(header.add_run("     /     2026 作品介绍"), size=8.5, color=MUTED)
    header._p.get_or_add_pPr().append(xml("w:pBdr"))
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(footer.add_run("MemHall Team    ·    "), size=8, color=MUTED)
    field(footer, "PAGE")
    # Footer page fields must not carry the TOC placeholder text.
    for run in footer.runs:
        if run.text == "目录将在 Word 中更新":
            run.text = "1"
    doc.core_properties.title = "麟阁 MemHall · 作品介绍"
    doc.core_properties.subject = "2026 上海开源软件应用创新大赛 · openKylin 智能体长期记忆自动化评测"
    doc.core_properties.author = "MemHall Team"
    doc.core_properties.keywords = "openKylin, MemHall, agent memory, benchmark, open source"
    doc.core_properties.comments = (
        "Editable text and tables; historical results and audit limitations are explicitly disclosed."
    )


def make_figures(summary):
    ASSETS.mkdir(exist_ok=True)
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 10,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "savefig.dpi": 250,
            "axes.labelcolor": "#142C46",
        }
    )
    with (OUT / "experiment-rounds.csv").open(encoding="utf-8-sig", newline="") as stream:
        rounds = list(csv.DictReader(stream))
    assert len(rounds) == 18
    fig, ax = plt.subplots(figsize=(7.2, 3.2))
    for i, (agent, color) in enumerate(zip(ORDER, COLORS, strict=True)):
        vals = np.array(
            [
                float(x["overall_score"]) * 100
                for x in rounds
                if x["scope"] == "full" and x["adapter"] == agent
            ]
        )
        assert len(vals) == 4
        assert abs(vals.mean() / 100 - summary["full"][agent]["mean"]) < 1e-8
        ax.bar(i, vals.mean(), width=0.6, color=color, alpha=0.13, edgecolor=color, lw=1.4, zorder=2)
        ax.errorbar(
            i, vals.mean(), yerr=vals.std(ddof=1), fmt="none", ecolor=color, capsize=5, lw=1.7, zorder=4
        )
        ax.scatter(
            i + np.array([-0.12, -0.04, 0.04, 0.12]),
            vals,
            s=35,
            color=color,
            edgecolors="white",
            lw=0.7,
            zorder=5,
        )
        ax.text(
            i,
            94,
            f"{vals.mean():.1f}% ± {vals.std(ddof=1):.1f}",
            ha="center",
            color=color,
            fontsize=12,
            weight="bold",
        )
    ax.set(
        ylim=(0, 102),
        xlim=(-0.6, 2.6),
        ylabel="Overall accuracy (%)",
        xticks=range(3),
        xticklabels=["Hermes", "OpenClaw", "KylinBot"],
    )
    ax.set_yticks(range(0, 101, 20))
    ax.grid(axis="y", color="#DFE6EC", lw=0.6)
    ax.set_axisbelow(True)
    fig.tight_layout(pad=1.0)
    fig.savefig(ASSETS / "round-scores.png", bbox_inches="tight", facecolor="white")
    fig.savefig(ASSETS / "round-scores.svg", bbox_inches="tight", facecolor="white")
    plt.close(fig)
    # Layout and glyphs adapted from paper-figures case_seg057; data are actual frozen results.
    theta = np.deg2rad(90 - np.arange(6) * 60)
    fig = plt.figure(figsize=(8, 3.25))
    for i, (agent, color) in enumerate(zip(ORDER, COLORS, strict=True)):
        ax = fig.add_axes([0.015 + i * 0.332, 0.1, 0.315, 0.79])
        ax.set(xlim=(-1.5, 1.5), ylim=(-1.5, 1.5), aspect="equal")
        ax.axis("off")
        ax.add_patch(Wedge((0, 0), 1.17, 0, 180, width=0.10, fc="#D9ECEC", ec="none"))
        ax.add_patch(Wedge((0, 0), 1.17, 180, 360, width=0.10, fc="#E5EAF0", ec="none"))
        for r in [0.25, 0.5, 0.75, 1]:
            ax.add_patch(Circle((0, 0), r, fill=False, ec="#CBD6E0", lw=0.65))
        for t in theta:
            ax.plot([0, np.cos(t)], [0, np.sin(t)], color="#DEE5EC", lw=0.65)
        vals = np.array([summary["full"][agent]["caps"][key] for key in CAPS])
        x, y = vals * np.cos(theta), vals * np.sin(theta)
        ax.fill(x, y, color=color, alpha=0.26)
        ax.plot(np.r_[x, x[0]], np.r_[y, y[0]], color=color, lw=1.7)
        ax.scatter(x, y, s=22, color=color, zorder=3)
        for label, t in zip("ABCDEF", theta, strict=True):
            ax.text(
                1.34 * np.cos(t), 1.34 * np.sin(t), label, ha="center", va="center", size=12, color="#142C46"
            )
        ax.text(0.02, 0.5, "50", size=7, color="#718296")
        ax.text(0.02, 1, "100", size=7, color="#718296")
        fig.text(
            0.173 + i * 0.332,
            0.95,
            ["Hermes", "OpenClaw", "KylinBot"][i],
            color=color,
            ha="center",
            weight="bold",
            size=12,
        )
    fig.savefig(ASSETS / "radar-profiles.png", bbox_inches="tight", facecolor="white")
    fig.savefig(ASSETS / "radar-profiles.svg", bbox_inches="tight", facecolor="white")
    plt.close(fig)


def results(doc, summary, scope):
    if scope == "full":
        headers = ["智能体", "总体 mean±SD", "保持", "调用", "更新", "区分", "边界", "复用"]
        rows = []
        for agent in ORDER:
            record = summary[scope][agent]
            rows.append(
                [agent, f"{record['mean']:.1%} ± {record['sd'] * 100:.1f}"]
                + [f"{record['caps'][key]:.1%}" for key in CAPS]
            )
        editable_table(doc, headers, rows, [2.0, 3.1] + [11.5 / 6] * 6)
    else:
        rows = [
            [
                agent,
                f"{summary[scope][agent]['mean']:.1%} ± {summary[scope][agent]['sd'] * 100:.1f}",
                f"{summary[scope][agent]['caps']['persist']:.1%}",
                f"{summary[scope][agent]['caps']['reuse']:.1%}",
            ]
            for agent in ORDER
        ]
        editable_table(
            doc, ["智能体", "任务链总体 mean±SD", "长期保持", "任务复用"], rows, [3.5, 5.1, 4.0, 4.0]
        )


def build():
    content = json.loads((OUT / "submission-content.json").read_text(encoding="utf-8"))
    summary = json.loads((OUT / "experiment-summary.json").read_text(encoding="utf-8"))
    make_figures(summary)
    doc = Document()
    configure(doc)
    paragraph(doc, "2026 上海开源软件应用创新大赛", size=12, color=TEAL, bold=True)
    paragraph(doc, "开源 AI 工具赛道 · openKylin 企业命题", size=10, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(20)
    cover = doc.add_table(rows=1, cols=1)
    cell = cover.cell(0, 0)
    shade(cell, INK)
    para = cell.paragraphs[0]
    para.paragraph_format.space_before = Pt(35)
    para.paragraph_format.space_after = Pt(15)
    font(para.add_run("麟阁"), size=36, bold=True, color="FFFFFF")
    para = paragraph(cell, "MEMHALL", size=29, bold=True, color="7BDFD5")
    para.paragraph_format.space_after = Pt(25)
    paragraph(cell, "作品介绍", size=19, color="FFFFFF", bold=True)
    paragraph(cell, "面向 openKylin 的智能体\n长期记忆自动化评测基准", size=16, color="FFFFFF")
    para = paragraph(cell, "教 — 隔 — 考   /   四类证据   /   六维能力", size=11, color="C9DAE7")
    para.paragraph_format.space_after = Pt(35)
    paragraph(doc, "让记忆能力，经得起复现。", size=18, color=TEAL, bold=True)
    paragraph(doc, "统一实验流程 · 可解释判卷 · 系统级验证 · 公开测量边界", size=10, color=MUTED)
    editable_table(
        doc,
        ["49 道公开基准题", "6 项能力维度", "3 款真实智能体"],
        [["45 full + 4 chains", "保持 / 调用 / 更新 / 区分 / 边界 / 复用", "full 各 4 轮 · chains 各 2 轮"]],
        [5.3, 6.0, 5.3],
    )
    paragraph(doc, "MemHall Team  |  2026.10.13", size=10, color=MUTED)
    paragraph(doc, "可编辑评审稿 · 基于 v1.3.2 源码与逐轮实验留档", size=9, color=MUTED)
    hyperlink(doc.add_paragraph(), "开源代码仓库", "https://github.com/shangjian2023/MemHall")
    nav = paragraph(doc, "阅读导航", size=19, bold=True, color=INK)
    nav.paragraph_format.page_break_before = True
    paragraph(doc, "评审可先看作品速览、技术架构、真实实验与创新点，再沿证据及复现章节核验。", color=MUTED)
    field(doc.add_paragraph(), r'TOC \o "1-1" \h \z \u')
    paragraph(
        doc,
        "目录为 Word 自动域；可在 Word 中右键更新域。正文及全部表格支持直接修改，图表数据另附 CSV。",
        size=9,
        color=MUTED,
    )
    markdown = ["# 麟阁 MemHall · 作品介绍", "", content["subtitle"], "", "核验日期：2026-10-13", ""]
    for page in content["pages"]:
        chapter = paragraph(doc, page["title"], style="Heading 1")
        chapter.paragraph_format.page_break_before = True
        paragraph(doc, page["deck"], color=TEAL, bold=True, size=10.6)
        markdown += ["## " + page["title"], "", page["deck"], ""]
        for block in page["blocks"]:
            kind = block["type"]
            if kind == "p":
                paragraph(doc, block["text"])
                markdown += [block["text"], ""]
            elif kind == "h":
                paragraph(doc, block["text"], style="Heading 2")
                markdown += ["### " + block["text"], ""]
            elif kind == "note":
                para = paragraph(doc, block["text"], size=9.1, color=MUTED)
                para._p.get_or_add_pPr().append(xml("w:shd", **{"w:fill": LIGHT}))
                markdown += ["> " + block["text"], ""]
            elif kind == "table":
                editable_table(doc, block["headers"], block["rows"])
                markdown += [
                    "| " + " | ".join(block["headers"]) + " |",
                    "| " + " | ".join(["---"] * len(block["headers"])) + " |",
                ]
                markdown += ["| " + " | ".join(row) + " |" for row in block["rows"]] + [""]
            elif kind == "image":
                file = ASSETS / block["name"]
                if file.is_file():
                    doc.add_picture(str(file), width=Cm(block.get("width", 16)))
                    paragraph(doc, block["caption"], style="Caption")
                    markdown += [f"![{block['caption']}](assets/{file.name})", ""]
                else:
                    raise FileNotFoundError(file)
            elif kind == "code":
                para = paragraph(doc, block["text"], size=8.7)
                para.paragraph_format.line_spacing = 1.12
                para._p.get_or_add_pPr().append(xml("w:shd", **{"w:fill": LIGHT}))
                for run in para.runs:
                    run.font.name = "Consolas"
                markdown += ["```bash", block["text"], "```", ""]
            elif kind.startswith("results_"):
                results(doc, summary, kind.removeprefix("results_"))
                markdown += ["数据见 experiment-summary.json 与 experiment-rounds.csv。", ""]
    target = OUT / "MemHall-作品介绍-可编辑版.docx"
    doc.save(target)
    (OUT / "MemHall-作品介绍.md").write_text("\n".join(markdown), encoding="utf-8")
    print(f"Built {target} ({target.stat().st_size:,} bytes); {len(content['pages'])} content sections.")
    print(f"Frozen-data figure validation: {len(ORDER)} groups × 4 observed rounds, no fabricated data.")


if __name__ == "__main__":
    build()
