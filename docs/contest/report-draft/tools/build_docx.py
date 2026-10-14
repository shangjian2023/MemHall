# -*- coding: utf-8 -*-
"""报告 markdown 方言 → Word 渲染器（格式契约见 ../format-spec.md）。

用法：
    python build_docx.py --module src/ch03-overview.md   # 单模块 → out/ch03-overview.docx
    python build_docx.py --all                           # 封面+目录+全部模块+页码 → out/麟阁参赛报告-整本.docx

图表编号：按各语义 ID 在全部模块（文件名排序）中的出现顺序分配连续图号/表号；
单模块渲染也用全书编号，保证与整本一致。正文引用 `图~id~` / `表~id~` 渲染时替换。
"""
import argparse
import glob
import io
import os
import re
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8")
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(BASE, "src")
FIGDIR = os.path.join(BASE, "figures")
OUT = os.path.join(BASE, "out")

EA_BODY, EA_HEAD, EA_CAP = "宋体", "黑体", "楷体"
ASCII_FONT = "Times New Roman"
MONO = "Consolas"
GRAY = RGBColor(0x80, 0x80, 0x80)

FIG_RE = re.compile(r"^\{\{fig:([^|}]+)\|([^|}]+)\|([^}]+)\}\}\s*$")
TAB_RE = re.compile(r"^【表:([^|】]+)\|([^】]+)】\s*$")
REF_FIG = re.compile(r"图~([A-Za-z0-9:_.\-]+)~")
REF_TAB = re.compile(r"表~([A-Za-z0-9:_.\-]+)~")
INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|〔待填[:：][^〕]+〕|图~[A-Za-z0-9:_.\-]+~|表~[A-Za-z0-9:_.\-]+~)")
MD_LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def set_font(run, ea=EA_BODY, ascii_=ASCII_FONT, size=12, bold=None, color=None):
    run.font.name = ascii_
    if run._element.rPr is not None and run._element.rPr.rFonts is not None:
        run._element.rPr.rFonts.set(qn("w:eastAsia"), ea)
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade(run_or_p_el, fill="F2F2F2"):
    el = run_or_p_el._element if hasattr(run_or_p_el, "_element") else run_or_p_el
    pr = el.get_or_add_pPr() if el.tag == qn("w:p") else el.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def p_border(par, top=True, bottom=True):
    pPr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    for side, on in (("top", top), ("bottom", bottom)):
        if on:
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "6")
            e.set(qn("w:color"), "808080")
            bdr.append(e)
    pPr.append(bdr)


def add_runs(par, text, figs, tabs, size=12, ea=EA_BODY):
    text = MD_LINK.sub(lambda m: f"{m.group(1)}（{m.group(2)}）", text)
    for tok in INLINE_RE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(tok[2:-2])
            set_font(r, ea=ea, size=size, bold=True)
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1])
            set_font(r, ea=ea, ascii_=MONO, size=size - 1.5)
            shade(r)
        elif tok.startswith("〔待填"):
            r = par.add_run(tok)
            set_font(r, ea=ea, size=size, color=GRAY)
            shade(r, fill="FFF2CC")
        elif REF_FIG.fullmatch(tok) or REF_TAB.fullmatch(tok):
            m = REF_FIG.match(tok) or REF_TAB.match(tok)
            kind = "figs" if tok.startswith("图") else "tabs"
            table = figs if kind == "figs" else tabs
            key = m.group(1)
            if key not in table:
                raise SystemExit(f"[编号缺失] {tok} 引用了未登记的 ID：{key}（检查 {{fig:…}}/【表:…】）")
            label = "图" if kind == "figs" else "表"
            r = par.add_run(f"{label}{table[key]}")
            set_font(r, ea=ea, size=size)
        else:
            r = par.add_run(tok)
            set_font(r, ea=ea, size=size)


def body_par(doc, text, figs, tabs, indent=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        pf.first_line_indent = Cm(0.85)
    add_runs(p, text, figs, tabs)
    return p


def quote_par(doc, text, figs, tabs):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.left_indent = Cm(0.74)
    pf.space_before = Pt(3)
    pf.space_after = Pt(3)
    add_runs(p, text, figs, tabs, size=11, ea=EA_CAP)
    return p


def list_par(doc, text, figs, tabs, marker):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.left_indent = Cm(0.85)
    add_runs(p, f"{marker} {text}", figs, tabs)
    return p


def caption_par(doc, label, title):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.line_spacing = 1.2
    pf.space_before = Pt(3)
    pf.space_after = Pt(6)
    r = p.add_run(f"{label} {title}")
    set_font(r, ea=EA_CAP, size=11)
    return p


def figure_block(doc, m, figs):
    fid, title, status = m.group(1).strip(), m.group(2).strip(), m.group(3).strip()
    png = os.path.join(FIGDIR, f"{fid}.png")
    if os.path.exists(png):
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        run = p.add_run()
        run.add_picture(png, width=Cm(14))
    else:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.space_before = Pt(10)
        pf.space_after = Pt(10)
        r = p.add_run(f"［此处插入图：{title}（{status}）］")
        set_font(r, ea=EA_CAP, size=11, color=GRAY)
        p_border(p)
    caption_par(doc, f"图{figs[fid]}", title)


def set_cell(cell, text, figs, tabs, header=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.15
    add_runs(p, text.strip(), figs, tabs, size=11)
    if header:
        for r in p.runs:
            r.font.bold = True


def render_table(doc, rows, figs, tabs):
    header, body = rows[0], [r for r in rows[1:] if not re.match(r"^\s*\|?[\s:\-|]+\|?\s*$", r)]
    cols = [c.strip() for c in header.strip().strip("|").split("|")]
    t = doc.add_table(rows=1, cols=len(cols))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side, sz in (("top", "12"), ("bottom", "12")):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), sz)
        e.set(qn("w:color"), "000000")
        borders.append(e)
    tblPr.append(borders)
    for i, name in enumerate(cols):
        set_cell(t.rows[0].cells[i], name, figs, tabs, header=True)
        tcPr = t.rows[0].cells[i]._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        e = OxmlElement("w:bottom")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), "000000")
        b.append(e)
        tcPr.append(b)
    for line in body:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        row = t.add_row()
        for i in range(len(cols)):
            set_cell(row.cells[i], cells[i] if i < len(cells) else "", figs, tabs)


def parse_lines(text):
    """把 markdown 源切成 (kind, payload) 行流。"""
    out = []
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            buf = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            out.append(("code", buf))
            continue
        if s.startswith("|"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                buf.append(lines[i])
                i += 1
            out.append(("table", buf))
            continue
        if s.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            out.append(("quote", " ".join(x for x in buf if x)))
            continue
        if s.startswith("#"):
            out.append(("h", s))
        elif s.startswith("- "):
            out.append(("li", ("•", s[2:].strip())))
        elif re.match(r"^\d+[.、]\s*", s):
            out.append(("li", (re.match(r"^(\d+)[.、]\s*", s).group(1) + ".", s[re.match(r"^\d+[.、]\s*", s).end():].strip())))
        else:
            out.append(("p", s))
        i += 1
    return out


def numbering_pass(module_files):
    figs, tabs = {}, {}
    fn, tn = 0, 0
    for path in module_files:
        text = open(path, encoding="utf-8").read()
        for kind, payload in parse_lines(text):
            if kind == "table" and len(payload) > 1:
                continue
            src = payload if isinstance(payload, str) else str(payload)
            for m in FIG_RE.finditer("\n".join(payload) if isinstance(payload, list) else src):
                if m.group(1).strip() not in figs:
                    fn += 1
                    figs[m.group(1).strip()] = fn
            for m in TAB_RE.finditer("\n".join(payload) if isinstance(payload, list) else src):
                if m.group(1).strip() not in tabs:
                    tn += 1
                    tabs[m.group(1).strip()] = tn
    return figs, tabs


def style_setup(doc):
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.top_margin = sec.bottom_margin = Cm(2.54)
        sec.left_margin = Cm(2.86)
        sec.right_margin = Cm(2.59)
        sec.header_distance, sec.footer_distance = Cm(1.5), Cm(1.75)
    st = doc.styles["Normal"]
    st.font.name = ASCII_FONT
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), EA_BODY)
    specs = {
        "Heading 1": (EA_HEAD, 16, True, WD_ALIGN_PARAGRAPH.CENTER, 17, 16.5),
        "Heading 2": (EA_BODY, 12, True, None, 13, 13),
        "Heading 3": (EA_BODY, 12, True, None, 6, 6),
    }
    for name, (ea, size, bold, align, before, after) in specs.items():
        h = doc.styles[name]
        h.font.name = ASCII_FONT
        h.font.size = Pt(size)
        h.font.bold = bold
        h.element.get_or_add_rPr()
        rf = h.element.rPr.get_or_add_rFonts()
        rf.set(qn("w:eastAsia"), ea)
        rf.set(qn("w:ascii"), ASCII_FONT)
        rf.set(qn("w:hAnsi"), ASCII_FONT)
        pf = h.paragraph_format
        if align:
            pf.alignment = align
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.keep_with_next = True


def heading(doc, text, figs, tabs):
    m = re.match(r"^(#{1,3})\s*(.+)$", text)
    level = len(m.group(1))
    p = doc.add_heading("", level=level)
    add_runs(p, m.group(2).strip(), figs, tabs, size=16 if level == 1 else 12,
             ea=EA_HEAD if level == 1 else EA_BODY)
    if level == 1:
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            set_font(r, ea=EA_HEAD, size=16, bold=True)
    else:
        for r in p.runs:
            set_font(r, ea=EA_BODY, size=12, bold=True)
    return p


def render_module(doc, path, figs, tabs):
    pending_tab = None
    for kind, payload in parse_lines(open(path, encoding="utf-8").read()):
        if kind == "h":
            heading(doc, payload, figs, tabs)
        elif kind == "p":
            m = FIG_RE.match(payload)
            if m:
                figure_block(doc, m, figs)
                continue
            m = TAB_RE.match(payload)
            if m:
                pending_tab = m
                continue
            body_par(doc, payload, figs, tabs)
        elif kind == "li":
            list_par(doc, payload[1], figs, tabs, payload[0])
        elif kind == "quote":
            quote_par(doc, payload, figs, tabs)
        elif kind == "table":
            rows = payload
            if rows and len(rows) >= 2:
                if pending_tab is not None:
                    caption_par(doc, f"表{tabs[pending_tab.group(1).strip()]}", pending_tab.group(2).strip())
                    pending_tab = None
                render_table(doc, rows, figs, tabs)
                doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif kind == "code":
            for cl in payload:
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.left_indent = Cm(0.85)
                pf.line_spacing = 1.15
                r = p.add_run(cl)
                set_font(r, ea=EA_BODY, ascii_=MONO, size=10.5)


def field(par, instr):
    r = par.add_run()
    f1, f2, f3 = OxmlElement("w:fldChar"), OxmlElement("w:instrText"), OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    f2.set(qn("xml:space"), "preserve")
    f2.text = instr
    f3.set(qn("w:fldCharType"), "end")
    r._r.append(f1)
    r._r.append(f2)
    r._r.append(f3)
    set_font(r, size=9)


def add_footer_pagenum(doc):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    p = sec.footer.paragraphs[0]
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(p, "PAGE")


def cover(doc):
    for line, size, before in ((f"openKylin 开源贡献赛 · 开源 AI 工具赛道", 26, 80),
                               ("作品报告", 26, 20)):
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before)
        r = p.add_run(line)
        set_font(r, ea=EA_HEAD, size=size, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(140)
    r = p.add_run("作品名称：麟阁 MemHall —— 面向 openKylin 的智能体长期记忆自动化评测基准")
    set_font(r, ea=EA_HEAD, size=16, bold=True)
    doc.add_page_break()


def toc(doc):
    p = doc.add_paragraph("目  录")
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        set_font(r, ea=EA_HEAD, size=16, bold=True)
    field(doc.add_paragraph(), r'TOC \o "1-3" \h \z \u')
    doc.add_page_break()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--module")
    ap.add_argument("--all", action="store_true")
    args = ap.parse_args()
    modules = sorted(p for p in glob.glob(os.path.join(SRC, "ch*.md")) if "selftest" not in p)
    figs, tabs = numbering_pass(modules)
    doc = Document()
    style_setup(doc)
    if args.all:
        abstract = os.path.join(SRC, "ch00-abstract.md")
        cover(doc)
        if os.path.exists(abstract):
            render_module(doc, abstract, figs, tabs)
            doc.add_page_break()
        toc(doc)
        add_footer_pagenum(doc)
        targets = [m for m in modules if "ch00" not in m]
    elif args.module:
        targets = [args.module]
    else:
        ap.error("需要 --module 或 --all")
    for path in targets:
        render_module(doc, path, figs, tabs)
    os.makedirs(OUT, exist_ok=True)
    if args.all:
        out = os.path.join(OUT, "麟阁参赛报告-整本.docx")
    else:
        out = os.path.join(OUT, os.path.basename(args.module).replace(".md", ".docx"))
    doc.save(out)
    n_pending = sum(1 for f in figs if not os.path.exists(os.path.join(FIGDIR, f"{f}.png")))
    n_fill = sum(open(m, encoding="utf-8").read().count("〔待填")
                 for m in targets)
    print(f"OK {out}")
    print(f"  图 {len(figs)} 个（占位 {n_pending}）｜表 {len(tabs)} 张｜待填标记 {n_fill} 处")


if __name__ == "__main__":
    main()
