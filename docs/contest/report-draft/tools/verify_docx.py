# -*- coding: utf-8 -*-
"""整本 docx 逐 run 核验：颜色全黑（灰占位除外）、各样式字体字号与 format-spec 一致。"""
import io
import sys
from collections import Counter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
from docx import Document
from docx.oxml.ns import qn

doc = Document(sys.argv[1] if len(sys.argv) > 1 else
               "docs/contest/report-draft/out/麟阁参赛报告-整本.docx")

ALLOWED = {"000000", "808080"}          # 黑 + 图占位灰
SPEC = {  # 样式 → (eastAsia 允许集, 字号允许集, 允许加粗)
    "Heading 1": ({"黑体"}, {16.0}, {True}),
    "Heading 2": ({"宋体"}, {12.0}, {True}),
    "Heading 3": ({"宋体"}, {12.0}, {True}),
    "Normal": ({"宋体", "楷体", "黑体"}, {9.0, 10.5, 11.0, 12.0, 16.0, 26.0}, {None, True, False}),  # 16/26=目录与封面（builder 按规范生成）
}

bad, style_stats = [], Counter()
for p in doc.paragraphs:
    if not p.text.strip() and not p.runs:
        continue
    st = p.style.name
    for r in p.runs:
        if not r.text.strip():
            continue
        rpr = r._element.rPr
        ea = rpr.rFonts.get(qn("w:eastAsia")) if (rpr is not None and rpr.rFonts is not None) else None
        color = r.font.color.rgb
        color_hex = str(color) if color is not None else None
        size = r.font.size.pt if r.font.size else None
        bold = r.font.bold
        style_stats[(st, ea, size, bold)] += 1
        if color_hex is not None and color_hex not in ALLOWED:
            bad.append(f"[颜色] {st} color={color_hex} :: {r.text[:20]}")
        if st in SPEC:
            eas, sizes, bolds = SPEC[st]
            if ea not in eas:
                bad.append(f"[中文字体] {st} eastAsia={ea} :: {r.text[:20]}")
            allowed_sizes = sizes | ({16.0} if r.text.strip() == "目  录" else set()) \
                | ({9.5} if (r.font.name == "Consolas") else set())
            if size not in allowed_sizes:
                bad.append(f"[字号] {st} {size}pt :: {r.text[:20]}")
            if bold not in bolds:
                bad.append(f"[加粗] {st} bold={bold} :: {r.text[:20]}")

# 表格内文字同样扫一遍
for ti, t in enumerate(doc.tables):
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    if not r.text.strip():
                        continue
                    color = r.font.color.rgb
                    if color is not None and str(color) not in ALLOWED:
                        bad.append(f"[表{ti}颜色] {str(color)} :: {r.text[:16]}")
                    size = r.font.size.pt if r.font.size else None
                    if size not in (9.5, 10.5, 11.0, 12.0):  # 9.5 = 表内行内代码
                        bad.append(f"[表{ti}字号] {size}pt :: {r.text[:16]}")

print("== 样式×字体×字号×加粗 分布（段落） ==")
for (st, ea, size, bold), n in sorted(style_stats.items(), key=str):
    print(f"  {st:<10} eastAsia={ea or '-':<4} {size}pt bold={bold} ×{n}")
print("\n== 违规明细 ==" if bad else "\n== 零违规：颜色全黑（灰=图占位），字体字号加粗全部符合 spec ==")
for b in bad[:40]:
    print(" ", b)
print(f"\nparagraphs={len(doc.paragraphs)} tables={len(doc.tables)} bad={len(bad)}")
