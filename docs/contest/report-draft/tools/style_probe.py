# -*- coding: utf-8 -*-
"""从获奖报告 docx 提取排版规范：页面、样式、字体、字号、行距、图注表注、表格样式。"""
import io
import re
import sys
from collections import Counter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm


def cm(v):
    try:
        return round(Cm(v).cm, 2) if v is not None else None
    except Exception:
        return None


def eastasia_of(font_el):
    rpr = font_el.element.rPr if font_el.element is not None else None
    if rpr is None:
        return None
    rfonts = rpr.rFonts
    if rfonts is None:
        return None
    return rfonts.get(qn("w:eastAsia"))


def run_info(run):
    rpr = run._element.rPr
    ea = None
    if rpr is not None and rpr.rFonts is not None:
        ea = rpr.rFonts.get(qn("w:eastAsia"))
    return {
        "text": run.text[:24],
        "ascii": run.font.name,
        "eastAsia": ea,
        "size_pt": run.font.size.pt if run.font.size else None,
        "bold": run.font.bold,
    }


def para_info(p):
    pf = p.paragraph_format
    runs = [run_info(r) for r in p.runs if r.text.strip()][:2]
    return {
        "style": p.style.name,
        "align": str(p.alignment),
        "first_line_cm": cm(pf.first_line_indent) if pf.first_line_indent else None,
        "line_spacing": pf.line_spacing,
        "rule": str(pf.line_spacing_rule),
        "space_before": pf.space_before,
        "space_after": pf.space_after,
        "runs": runs,
    }


def main(path):
    doc = Document(path)
    print(f"===== {path} =====")
    for i, sec in enumerate(doc.sections):
        print(f"[section {i}] page {cm(sec.page_width)}x{cm(sec.page_height)}cm "
              f"margins T{cm(sec.top_margin)} B{cm(sec.bottom_margin)} "
              f"L{cm(sec.left_margin)} R{cm(sec.right_margin)} "
              f"header {cm(sec.header_distance)} footer {cm(sec.footer_distance)}")

    used = Counter(p.style.name for p in doc.paragraphs)
    print("\n[used paragraph styles]", dict(used))

    print("\n[style definitions]")
    for name in sorted(used) + ["Normal", "Title"]:
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        f = st.font
        ea = eastasia_of(f)
        pf = getattr(st, "paragraph_format", None)
        ls = pf.line_spacing if pf else None
        fi = cm(pf.first_line_indent) if pf and pf.first_line_indent else None
        sb = pf.space_before.pt if pf and pf.space_before else None
        sa = pf.space_after.pt if pf and pf.space_after else None
        print(f"  {name!r}: ascii={f.name} eastAsia={ea} size={f.size.pt if f.size else None} "
              f"bold={f.bold} line_spacing={ls} first_indent_cm={fi} before={sb} after={sa}")

    cap_re = re.compile(r"^\s*(图|表|Fig|Table)\s*[\d\-–—\.]")
    print("\n[caption paragraphs]")
    n = 0
    for p in doc.paragraphs:
        if cap_re.match(p.text):
            n += 1
            if n <= 12:
                info = para_info(p)
                print(f"  {info['style']} align={info['align']} ls={info['line_spacing']} "
                      f"runs={info['runs']} :: {p.text[:40]}")
    print(f"  ...total caption-like: {n}")

    print("\n[body sample: first 3 non-empty Normal paragraphs]")
    k = 0
    for p in doc.paragraphs:
        if p.text.strip() and p.style.name in ("Normal", "Body Text", "正文"):
            k += 1
            info = para_info(p)
            print(f"  align={info['align']} ls={info['line_spacing']} fi={info['first_line_cm']} "
                  f"runs={info['runs']} :: {p.text[:30]}")
            if k >= 3:
                break

    print(f"\n[tables] count={len(doc.tables)}")
    tstyles = Counter(t.style.name if t.style else None for t in doc.tables)
    print("  table styles:", dict(tstyles))
    for t in doc.tables[:2]:
        hdr = t.rows[0]
        cells = [c.text[:10] for c in hdr.cells][:4]
        run0 = hdr.cells[0].paragraphs[0].runs
        print(f"  first table {len(t.rows)}x{len(t.columns)} header={cells} "
              f"hdr_runs={run_info(run0[0]) if run0 else None}")
    print(f"\n[images] inline_shapes={len(doc.inline_shapes)} paragraphs={len(doc.paragraphs)}")


if __name__ == "__main__":
    main(sys.argv[1])
